from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference

wb = Workbook(write_only=True)
ws = wb.create_sheet()
#rows = [("None", "Item A","Item B"), (2013, 100, 50), (2014, 150, 100), (2015, 170, 120), (2016, 180, 130), (2017, 190, 140), (2018, 200, 150), (2019, 210, 170)]
rows = [("", "Item A","Item B"), (2013, 100, 50), (2014, 150, 100), (2015, 170, 120), (2016, 180, 130), (2017, 190, 140), (2018, 200, 150), (2019, 210, 170)]

for row in rows:
	ws.append(row)

c_line = LineChart()
c_line.title = "title"
c_line.style = 3
c_line.y_axis.title = "금액"
c_line.x_axis.title = "연도"
#                      열의 범위        ->|<-  data------------->
data1 = Reference(ws, min_col=2, max_col=3, min_row=1, max_row=8)
cats1 = Reference(ws, min_col=1, min_row=2, max_row=8)
c_line.add_data(data1, titles_from_data=True)
c_line.set_categories(cats1)
c_line.shape = 4
ws.add_chart(c_line, "A8")
wb.save("./data/line_chart.xlsx")
